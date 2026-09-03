"""
试卷导出工具：
- 用 markdown 库把题目 Markdown 渲染成 HTML（完美还原题目录入时的格式）
- Word 导出：HTML → htmldocx → python-docx
- PDF 导出：HTML → fpdf2.write_html

这样能正确处理：
  - 硬换行（题干、①②③④、ABCD 选项各占一行）
  - 表格（Markdown table → HTML table）
  - 图片（![alt](path) → <img>）
  - 段落、加粗、斜体等
"""

import os
import re
import io
from typing import List, Tuple, Optional

import markdown


# ============================================================
# Markdown → HTML
# ============================================================

# Markdown 表格分隔行正则：| --- | --- |
# 以 | 开头，只含 |、空格、-、:，以 | 结尾
_TABLE_SEP_PAT = re.compile(r"^\s*\|[\s\-:|]+\|\s*$")


def _is_table_sep(line: str) -> bool:
    """判断一行是否是 Markdown 表格分隔行（| --- | --- |）。
    分隔行必须含至少一个 -，否则纯空格的数据行 |  |  | 也会被误判。
    """
    s = line.strip()
    if not s.startswith("|") or "-" not in s:
        return False
    return _TABLE_SEP_PAT.match(line) is not None


def _preprocess_markdown(text: str) -> str:
    """
    预处理题目 Markdown：
    1. 在 Markdown 表格的表头行前插入空行，确保 markdown 库能识别表格
       （题目录入时表格前往往没空行，导致表格被当成普通文本）。
    2. 把单个 \\n（行内换行）转成 Markdown 硬换行（行尾两空格），
       这样 markdown 库会生成 <br>，保留题干/选项的换行。
       表格行（以 | 开头）不加硬换行，否则破坏表格。
    """
    if not text:
        return ""
    lines = text.split("\n")

    # 1. 在表格表头行前插入空行
    inserted = []
    for i, line in enumerate(lines):
        if (
            i + 1 < len(lines)
            and line.strip().startswith("|")
            and _is_table_sep(lines[i + 1])
            and (len(inserted) == 0 or inserted[-1].strip() != "")
        ):
            inserted.append("")  # 插入空行
        inserted.append(line)
    lines = inserted

    # 2. 硬换行处理
    processed = []
    for i, line in enumerate(lines):
        if i < len(lines) - 1 and lines[i + 1].strip() != "" and line.strip() != "":
            if line.strip().startswith("|"):
                processed.append(line)
            else:
                processed.append(line + "  ")
        else:
            processed.append(line)
    return "\n".join(processed)


def markdown_to_html(text: str) -> str:
    """把题目 Markdown 渲染成 HTML。"""
    if not text:
        return ""
    processed = _preprocess_markdown(text)
    return markdown.markdown(processed, extensions=["tables", "fenced_code"])


# ============================================================
# 图片处理：缺失图片的容错 + 限制宽度
# ============================================================

def _resolve_image_path_in_html(html: str, max_width_px: int = 680) -> str:
    """
    处理 HTML 里的 <img src="path">：
    - 缺失图片 → 替换成提示文本
    - 存在图片 → 转成绝对路径，并加 width 属性限制宽度（像素，供 htmldocx 用）
      max_width_px: 图片最大宽度（像素），默认 680px（约对应 A4 可用宽度）
    """
    from PIL import Image

    def replace_img(match):
        src = match.group(1)
        if not src:
            return match.group(0)
        abs_path = None
        if os.path.isabs(src) and os.path.exists(src):
            abs_path = src
        elif os.path.exists(src):
            abs_path = os.path.abspath(src)
        if abs_path is None:
            return f'<p>[图片缺失：{src}]</p>'
        try:
            with Image.open(abs_path) as im:
                w_px, h_px = im.size
            if w_px > max_width_px:
                w_px = max_width_px
            return f'<img src="{abs_path}" width="{w_px}" />'
        except Exception:
            return f'<p>[图片加载失败：{src}]</p>'

    return re.sub(r'<img[^>]*src=["\']([^"\']*)["\'][^>]*/?>', replace_img, html)


# ============================================================
# Word 导出（markdown → html → htmldocx）
# ============================================================

def export_to_word(
    questions: List[dict],
    paper_title: str,
    include_answer: bool = True,
    include_analysis: bool = True,
    include_source: bool = True,
) -> bytes:
    """把一套试卷的题目导出为 Word 文档，返回二进制内容。"""
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from htmldocx import HtmlToDocx

    doc = Document()

    # 设置默认中文字体（宋体）
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")

    parser = HtmlToDocx()

    def add_html_safe(html: str):
        """安全地把 HTML 加到 Word 文档，htmldocx 出错时降级为纯文本。"""
        if not html.strip():
            return
        try:
            parser.add_html_to_document(html, doc)
            # htmldocx 的 set_initial_attrs 不会重置 self.run，当 HTML 以 <table>
            # 结尾时 self.run 会残留指向表格前的段落，污染后续 add_html_to_document
            # 调用（表现为上一题的文本重复出现在后续题目中）。手动清空以避免污染。
            parser.run = None
        except Exception:
            plain = re.sub(r'<[^>]+>', '', html)
            plain = plain.replace('&nbsp;', ' ').replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&')
            for line in plain.split('\n'):
                if line.strip():
                    doc.add_paragraph(line.strip())

    # 标题
    title_p = doc.add_heading(paper_title, level=0)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_run = info_p.add_run(f"共 {len(questions)} 道题")
    info_run.font.size = Pt(10)

    doc.add_paragraph()

    for idx, q in enumerate(questions, 1):
        qno = q.get("question_no") or str(idx)
        content = q.get("content") or ""
        answer = q.get("answer") or ""
        analysis = q.get("analysis") or ""
        source = q.get("source") or ""
        analysis_source = q.get("analysis_source") or ""

        doc.add_heading(f"第 {qno} 题", level=1)

        html = markdown_to_html(content)
        html = _resolve_image_path_in_html(html)
        add_html_safe(html)

        if include_answer and answer:
            doc.add_heading("【答案】", level=2)
            html = markdown_to_html(answer)
            html = _resolve_image_path_in_html(html)
            add_html_safe(html)

        if include_analysis and analysis:
            doc.add_heading("【解析】", level=2)
            html = markdown_to_html(analysis)
            html = _resolve_image_path_in_html(html)
            add_html_safe(html)

        if include_source:
            src_parts = []
            if source:
                src_parts.append(f"题目来源：{source}")
            if analysis_source and analysis_source != source:
                src_parts.append(f"解析来源：{analysis_source}")
            if src_parts:
                p = doc.add_paragraph()
                run = p.add_run("；".join(src_parts))
                run.font.size = Pt(9)
                run.italic = True

        if idx < len(questions):
            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ============================================================
# PDF 导出（markdown → html → fpdf2）
# ============================================================

_CJK_FONT_CANDIDATES = [
    ("PingFang", "/System/Library/Fonts/PingFang.ttc"),
    ("STHeiti", "/System/Library/Fonts/STHeiti Medium.ttc"),
    ("HiraginoSansGB", "/System/Library/Fonts/Hiragino Sans GB.ttc"),
    ("NotoSansCJK", "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc"),
    ("NotoSansCJK", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ("NotoSansCJK", "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
    ("WenQuanYi", "/usr/share/fonts/wenquanyi/wqy-microhei/wqy-microhei.ttc"),
]


def _find_cjk_font() -> Tuple[str, str]:
    for name, path in _CJK_FONT_CANDIDATES:
        if os.path.exists(path):
            return name, path
    raise RuntimeError("未找到可用的中文字体。请安装 Noto CJK 或 PingFang 等中文字体。")


def _render_html_to_pdf(pdf, html: str):
    """
    把 HTML 渲染到 PDF。
    图片单独用 pdf.image() 渲染（精确控制宽度），其余 HTML 用 write_html。
    """
    from PIL import Image

    if not html.strip():
        return

    img_pat = re.compile(r'<img[^>]*src=["\']([^"\']*)["\'][^>]*/?>')
    parts = []
    last_end = 0
    for m in img_pat.finditer(html):
        if m.start() > last_end:
            parts.append(("html", html[last_end:m.start()]))
        parts.append(("image", m.group(1)))
        last_end = m.end()
    if last_end < len(html):
        parts.append(("html", html[last_end:]))

    # 修复被截断的 <p> 标签
    fixed_parts = []
    for kind, payload in parts:
        if kind == "html":
            h = payload
            open_p = len(re.findall(r'<p[^>]*>', h))
            close_p = len(re.findall(r'</p>', h))
            if open_p > close_p:
                h = h + '</p>' * (open_p - close_p)
            elif close_p > open_p:
                h = '<p>' * (close_p - open_p) + h
            payload = h
        fixed_parts.append((kind, payload))
    parts = fixed_parts

    page_width = pdf.w - 2 * pdf.l_margin

    for kind, payload in parts:
        if kind == "html":
            h = payload.strip()
            if h:
                try:
                    pdf.write_html(h)
                except Exception:
                    plain = re.sub(r'<[^>]+>', '', h)
                    plain = plain.replace('&nbsp;', ' ').replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&')
                    if plain.strip():
                        pdf.multi_cell(0, 7, plain, new_x="LMARGIN", new_y="NEXT")
        elif kind == "image":
            src = payload
            if not os.path.exists(src):
                pdf.multi_cell(0, 7, f"[图片缺失：{src}]", new_x="LMARGIN", new_y="NEXT")
                continue
            try:
                with Image.open(src) as im:
                    w_px, h_px = im.size
                w_mm = w_px * 25.4 / 96.0
                if w_mm > page_width:
                    w_mm = page_width
                pdf.image(src, w=w_mm)
                pdf.ln(2)
            except Exception:
                pdf.multi_cell(0, 7, f"[图片加载失败：{src}]", new_x="LMARGIN", new_y="NEXT")


def export_to_pdf(
    questions: List[dict],
    paper_title: str,
    include_answer: bool = True,
    include_analysis: bool = True,
    include_source: bool = True,
) -> bytes:
    """把一套试卷的题目导出为 PDF，返回二进制内容。"""
    from fpdf import FPDF

    font_name, font_path = _find_cjk_font()

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.add_font(font_name, fname=font_path)
    pdf.add_font(font_name, style="B", fname=font_path)
    pdf.set_font(font_name, size=11)

    pdf.set_font_size(20)
    pdf.multi_cell(0, 12, paper_title, align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font_size(10)
    pdf.multi_cell(0, 8, f"共 {len(questions)} 道题", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    for idx, q in enumerate(questions, 1):
        qno = q.get("question_no") or str(idx)
        content = q.get("content") or ""
        answer = q.get("answer") or ""
        analysis = q.get("analysis") or ""
        source = q.get("source") or ""
        analysis_source = q.get("analysis_source") or ""

        pdf.set_font_size(14)
        pdf.multi_cell(0, 10, f"第 {qno} 题", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1)
        pdf.set_font_size(11)

        html = markdown_to_html(content)
        html = _resolve_image_path_in_html(html)
        _render_html_to_pdf(pdf, html)

        if include_answer and answer:
            pdf.ln(2)
            pdf.set_font_size(12)
            pdf.multi_cell(0, 9, "【答案】", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font_size(11)
            html = markdown_to_html(answer)
            html = _resolve_image_path_in_html(html)
            _render_html_to_pdf(pdf, html)

        if include_analysis and analysis:
            pdf.ln(2)
            pdf.set_font_size(12)
            pdf.multi_cell(0, 9, "【解析】", new_x="LMARGIN", new_y="NEXT")
            pdf.set_font_size(11)
            html = markdown_to_html(analysis)
            html = _resolve_image_path_in_html(html)
            _render_html_to_pdf(pdf, html)

        if include_source:
            src_parts = []
            if source:
                src_parts.append(f"题目来源：{source}")
            if analysis_source and analysis_source != source:
                src_parts.append(f"解析来源：{analysis_source}")
            if src_parts:
                pdf.ln(1)
                pdf.set_font_size(9)
                pdf.multi_cell(0, 6, "；".join(src_parts), new_x="LMARGIN", new_y="NEXT")
                pdf.set_font_size(11)

        pdf.ln(4)

    return bytes(pdf.output())
